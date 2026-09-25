"""Tests for the public historical export (/api/stations/<id>/export/).

Three things matter here and are covered separately:

* the **free-tier boundary** — the endpoint is unauthenticated, so every limit
  has to hold against a hand-written URL, not just against the picker that
  normally builds it;
* **eligibility** — only sensors Respira operates may be exported, checked on
  the server rather than trusted from the page;
* the **file itself** being well formed and carrying the right rows, which a
  200 alone cannot demonstrate.

Dates are relative to "today" rather than pinned to a calendar month, unlike
``tests_institution_exports``: the free-tier window is measured from the
server's clock, so a fixture pinned to July 2026 would start failing on its own
once that month fell out of range.
"""

import io
import json
from datetime import date, datetime, time, timedelta
from datetime import timezone as dt_timezone
from unittest.mock import patch

from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient, APITestCase
from rest_framework.throttling import ScopedRateThrottle

from .airgradient import AirGradientError, FetchResult
from .exports import REPORT_TIME_ZONE
from .models import Regions, StationDetails, Stations
from .public_exports import (
    FORMAT_PARAM,
    FREE_TIER_DAYS,
    MAX_PUBLIC_EXPORT_DAYS,
    ROW_COUNT_HEADER,
)


def _at(day: date, hour: int) -> datetime:
    """A measurement timestamp at a given local hour, as UTC."""
    return datetime.combine(day, time(hour), tzinfo=REPORT_TIME_ZONE).astimezone(
        dt_timezone.utc
    )


def _measure(moment: datetime, **overrides) -> dict:
    """One row shaped like the provider's own ``past`` payload.

    Mirrors the fixture in ``tests_institution_exports``, including the fields
    that arrive as strings: the export has to coerce them, and a fixture that
    handed it clean numbers would never exercise that.
    """
    row = {
        "locationId": 191355,
        "locationName": "Respira: Villa Morra",
        "pm01": 4.2,
        "pm02": 12.5,
        "pm10": 19.8,
        "rco2": 447,
        "atmp": 23.4,
        "rhum": 61.0,
        "tvocIndex": 102,
        "noxIndex": 1,
        "serialno": "588c813fe18c",
        "datapoints": "2",
        "timestamp": moment.astimezone(dt_timezone.utc).strftime(
            "%Y-%m-%dT%H:%M:%S.000Z"
        ),
    }
    row.update(overrides)
    return row


class PublicExportTestCase(APITestCase):
    def setUp(self):
        self.client = APIClient()
        # Throttling is a feature of this endpoint, so DRF's per-scope history
        # persists in the cache between tests and would make the second test in
        # a run fail on the first request. Cleared here; the throttle itself is
        # exercised deliberately in ThrottleTests.
        from django.core.cache import cache

        cache.clear()

        region = Regions.seed_for_tests(name="Gran Asuncion", region_code="GA")
        self.station = Stations.seed_for_tests(
            name="Respira: Villa Morra",
            # `respira_<locationId>` — eligibility is resolved from this.
            station_code="respira_191355",
            region=region,
            latitude=-25.28,
            longitude=-57.57,
            is_station_on=True,
        )
        StationDetails.objects.create(station=self.station, city="Asunción")

        self.foreign_station = Stations.seed_for_tests(
            name="FIUNA Campus",
            station_code="fiuna_01",
            region=region,
            latitude=-25.33,
            longitude=-57.51,
            is_station_on=True,
        )

        self.today = datetime.now(dt_timezone.utc).astimezone(REPORT_TIME_ZONE).date()
        # Three consecutive recent days, comfortably inside the free tier.
        self.days = [self.today - timedelta(days=offset) for offset in (3, 2, 1)]
        self.measures = [
            _measure(_at(day, hour)) for day in self.days for hour in (8, 20)
        ]

    def url(self, station=None) -> str:
        return reverse(
            "public-station-export",
            kwargs={"station_id": (station or self.station).id},
        )

    def stub(self, rows=None, failed_windows=0):
        """Stand in for the provider, capturing what was asked of it.

        Patches the import site in ``public_exports`` rather than the client
        module, and returns the captured call so a test can assert which sensor
        and which range were actually requested — the isolation guarantees the
        ticket asks for are about exactly that.
        """
        captured = {}

        def fake_fetch(location_id, start, end, **kwargs):
            captured["location_id"] = location_id
            captured["start"] = start
            captured["end"] = end
            return FetchResult(
                rows=list(self.measures if rows is None else rows),
                failed_windows=failed_windows,
            )

        patcher = patch(
            "api.public_exports.fetch_past_measures", side_effect=fake_fetch
        )
        patcher.start()
        self.addCleanup(patcher.stop)
        return captured

    def iso(self, day: date) -> str:
        return day.isoformat()

    def sheet(self, response):
        return load_workbook(io.BytesIO(response.content)).active

    def sheet_rows(self, response) -> list[tuple]:
        """Data rows of the spreadsheet, without the header."""
        return list(self.sheet(response).iter_rows(min_row=2, values_only=True))


class EligibilityTests(PublicExportTestCase):
    def test_anonymous_visitor_can_export(self):
        """The whole point of the free tier: no account, no contract."""
        self.stub()

        response = self.client.get(
            self.url(), {"from": self.iso(self.days[0]), "to": self.iso(self.days[-1])}
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(self.sheet_rows(response)), len(self.measures))

    def test_sensor_from_another_network_is_refused(self):
        """Eligibility holds on the server, not just in the page.

        The page hides the section for such a sensor, but hiding a control is
        not enforcement — a hand-written URL has to get the same answer.
        """
        self.stub()

        response = self.client.get(self.url(self.foreign_station))

        self.assertEqual(response.status_code, 404)

    def test_unknown_sensor_is_not_found(self):
        self.stub()

        response = self.client.get(
            reverse("public-station-export", kwargs={"station_id": 99999})
        )

        self.assertEqual(response.status_code, 404)

    def test_refusal_does_not_name_the_provider(self):
        """A public 404 explains the outcome without leaking the source.

        The panel's equivalent message may name the provider — it is read by
        staff. This one is read by anyone.
        """
        self.stub()

        response = self.client.get(self.url(self.foreign_station))

        self.assertNotIn("airgradient", str(response.json()).lower())


class FreeTierRangeTests(PublicExportTestCase):
    def test_start_before_the_free_tier_is_refused(self):
        self.stub()
        too_old = self.today - timedelta(days=FREE_TIER_DAYS + 5)

        response = self.client.get(
            self.url(), {"from": self.iso(too_old), "to": self.iso(self.today)}
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("from", response.json())
        # The limit is stated, so the caller can correct the request rather
        # than guess at it.
        self.assertIn(str(FREE_TIER_DAYS), str(response.json()))

    def test_oldest_allowed_day_is_served(self):
        """The boundary itself is inside the tier, not just short of it."""
        self.stub(rows=[])
        earliest = self.today - timedelta(days=FREE_TIER_DAYS - 1)

        response = self.client.get(
            self.url(), {"from": self.iso(earliest), "to": self.iso(self.today)}
        )

        self.assertEqual(response.status_code, 200)

    def test_inverted_range_is_refused(self):
        self.stub()

        response = self.client.get(
            self.url(),
            {
                "from": self.iso(self.today),
                "to": self.iso(self.today - timedelta(days=5)),
            },
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("to", response.json())

    def test_malformed_date_is_refused(self):
        self.stub()

        response = self.client.get(self.url(), {"from": "01-07-2026"})

        self.assertEqual(response.status_code, 400)

    def test_span_over_the_file_cap_is_refused(self):
        self.stub()

        with patch("api.public_exports.MAX_PUBLIC_EXPORT_DAYS", 2):
            response = self.client.get(
                self.url(),
                {
                    "from": self.iso(self.today - timedelta(days=5)),
                    "to": self.iso(self.today),
                },
            )

        self.assertEqual(response.status_code, 400)

    def test_future_end_date_cannot_widen_the_window(self):
        """A `to` past today is clamped, not treated as extra room.

        Without the clamp, `to=today+60` would let `from` sit 60 days earlier
        than the tier allows while still measuring a legal span — the window
        would have been slid forward instead of bounded.
        """
        captured = self.stub(rows=[])
        far_future = self.today + timedelta(days=60)

        response = self.client.get(
            self.url(),
            {
                "from": self.iso(self.today - timedelta(days=2)),
                "to": self.iso(far_future),
            },
        )

        self.assertEqual(response.status_code, 200)
        # The fetch stops at the end of today, not out in the future.
        self.assertLessEqual(
            captured["end"].astimezone(REPORT_TIME_ZONE).date(),
            self.today + timedelta(days=1),
        )

    def test_default_range_is_served_without_parameters(self):
        captured = self.stub()

        response = self.client.get(self.url())

        self.assertEqual(response.status_code, 200)
        requested_days = (captured["end"] - captured["start"]).days
        self.assertLessEqual(requested_days, MAX_PUBLIC_EXPORT_DAYS)


class ExportContentTests(PublicExportTestCase):
    def params(self) -> dict:
        return {"from": self.iso(self.days[0]), "to": self.iso(self.days[-1])}

    def test_spreadsheet_has_the_expected_columns(self):
        self.stub()

        response = self.client.get(self.url(), self.params())

        header = [cell.value for cell in self.sheet(response)[1]]
        self.assertEqual(
            header[:4],
            [
                "Sensor",
                "Sensor ID",
                "Fecha/Hora (local)",
                "Fecha/Hora (UTC)",
            ],
        )
        self.assertIn("PM2.5 (µg/m³)", header)
        self.assertIn("CO2 (ppm)", header)
        self.assertIn("Humedad (%)", header)

    def test_spreadsheet_identifies_the_sensor_by_our_own_id(self):
        """Never the upstream location id, which is not ours to publish."""
        self.stub()

        response = self.client.get(self.url(), self.params())
        first = self.sheet_rows(response)[0]

        self.assertEqual(first[0], "Respira: Villa Morra")
        self.assertEqual(first[1], self.station.id)

    def test_spreadsheet_rows_run_oldest_first(self):
        self.stub()

        response = self.client.get(self.url(), self.params())
        stamps = [row[3] for row in self.sheet_rows(response)]

        self.assertEqual(stamps, sorted(stamps))

    def test_spreadsheet_coerces_string_numbers(self):
        """The provider sends some values as text; a cell must hold a number."""
        self.stub(rows=[_measure(_at(self.days[0], 8), rco2="512")])

        response = self.client.get(self.url(), self.params())
        header = [cell.value for cell in self.sheet(response)[1]]
        co2 = self.sheet_rows(response)[0][header.index("CO2 (ppm)")]

        self.assertEqual(co2, 512)
        self.assertNotIsInstance(co2, str)

    def test_json_carries_the_measurements_and_its_context(self):
        self.stub()

        response = self.client.get(self.url(), {**self.params(), FORMAT_PARAM: "json"})
        payload = json.loads(response.content)

        self.assertEqual(payload["sensor"]["id"], self.station.id)
        self.assertEqual(payload["sensor"]["name"], "Respira: Villa Morra")
        self.assertEqual(payload["sensor"]["city"], "Asunción")
        self.assertEqual(payload["range"]["from"], self.iso(self.days[0]))
        self.assertEqual(payload["range"]["to"], self.iso(self.days[-1]))
        self.assertEqual(payload["measurement_count"], len(self.measures))
        self.assertEqual(len(payload["measurements"]), len(self.measures))

    def test_json_measurements_use_neutral_keys(self):
        """Our own vocabulary, not the provider's field names."""
        self.stub(rows=[_measure(_at(self.days[0], 8))])

        response = self.client.get(self.url(), {**self.params(), FORMAT_PARAM: "json"})
        measurement = json.loads(response.content)["measurements"][0]

        self.assertIn("pm2_5", measurement)
        self.assertIn("co2", measurement)
        self.assertIn("timestamp_utc", measurement)
        # The provider's own spellings must not leak into a published file.
        self.assertNotIn("pm02", measurement)
        self.assertNotIn("rco2", measurement)

    def test_unsupported_format_is_refused(self):
        self.stub()

        response = self.client.get(self.url(), {**self.params(), FORMAT_PARAM: "csv"})

        self.assertEqual(response.status_code, 400)
        self.assertIn(FORMAT_PARAM, response.json())

    def test_format_is_not_read_from_drf_s_reserved_parameter(self):
        """`?format=` belongs to DRF's content negotiation, not to us.

        Naming our parameter `format` made `?format=csv` resolve as a renderer
        lookup and answer 404 before this view ran at all — the caller asking
        for an unsupported format was told the sensor did not exist. Pinned
        here so the parameter is never renamed back.
        """
        self.stub()

        response = self.client.get(self.url(), {**self.params(), "format": "json"})

        # Untouched by the stray parameter: still the default spreadsheet.
        self.assertEqual(response.status_code, 200)
        self.assertIn(".xlsx", response["Content-Disposition"])

    def test_filename_names_the_sensor_and_the_range(self):
        self.stub()

        response = self.client.get(self.url(), self.params())

        disposition = response["Content-Disposition"]
        self.assertIn("respira-villa-morra", disposition)
        self.assertIn(self.iso(self.days[0]), disposition)
        self.assertIn(".xlsx", disposition)

    def test_filename_does_not_stutter_on_a_respira_named_sensor(self):
        """Most sensors are named "Respira: …"; the prefix must not double up."""
        self.stub()

        response = self.client.get(self.url(), self.params())

        self.assertNotIn("respira-respira", response["Content-Disposition"])

    def test_filename_does_not_name_the_provider(self):
        self.stub()

        response = self.client.get(self.url(), self.params())

        self.assertNotIn("airgradient", response["Content-Disposition"].lower())


class IsolationTests(PublicExportTestCase):
    def test_only_the_requested_sensor_is_fetched(self):
        eligible_other = Stations.seed_for_tests(
            name="Respira: Sajonia",
            station_code="respira_192812",
            is_station_on=True,
        )
        captured = self.stub()

        self.client.get(self.url(eligible_other))

        self.assertEqual(captured["location_id"], 192812)

    def test_only_the_requested_range_is_fetched(self):
        captured = self.stub()
        start, end = self.days[0], self.days[-1]

        self.client.get(self.url(), {"from": self.iso(start), "to": self.iso(end)})

        self.assertEqual(captured["start"].astimezone(REPORT_TIME_ZONE).date(), start)
        # The upper bound is exclusive: the day after the inclusive `to`.
        self.assertEqual(
            captured["end"].astimezone(REPORT_TIME_ZONE).date(),
            end + timedelta(days=1),
        )

    def test_measurements_outside_the_range_are_dropped(self):
        """A provider window may overshoot; the file must not."""
        inside = _measure(_at(self.days[-1], 10))
        outside = _measure(_at(self.today - timedelta(days=40), 10))
        self.stub(rows=[outside, inside])

        response = self.client.get(
            self.url(), {"from": self.iso(self.days[0]), "to": self.iso(self.days[-1])}
        )

        self.assertEqual(len(self.sheet_rows(response)), 1)


class EmptyAndFailureTests(PublicExportTestCase):
    def test_no_measurements_still_returns_a_file(self):
        """An empty period is a state, not an error.

        A 404 or a 204 would make the page guess whether the export failed; the
        row-count header says plainly that the sensor reported nothing.
        """
        self.stub(rows=[])

        response = self.client.get(self.url())

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response[ROW_COUNT_HEADER], "0")
        self.assertEqual(self.sheet_rows(response), [])

    def test_row_count_header_is_always_present(self):
        self.stub()

        response = self.client.get(self.url())

        self.assertEqual(response[ROW_COUNT_HEADER], str(len(self.measures)))

    def test_partial_export_is_served_and_flagged(self):
        self.stub(failed_windows=2)

        response = self.client.get(self.url())

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["X-Respira-Partial-Export"], "2")

    def test_complete_export_carries_no_partial_header(self):
        self.stub()

        response = self.client.get(self.url())

        self.assertNotIn("X-Respira-Partial-Export", response)

    def test_provider_failure_is_reported_without_naming_it(self):
        with patch(
            "api.public_exports.fetch_past_measures",
            side_effect=AirGradientError("no token"),
        ):
            response = self.client.get(self.url())

        self.assertEqual(response.status_code, 400)
        body = str(response.json()).lower()
        self.assertIn("detail", response.json())
        # The public must not learn who the upstream is, nor that a token is
        # what failed.
        self.assertNotIn("airgradient", body)
        self.assertNotIn("token", body)


class ThrottleTests(PublicExportTestCase):
    def test_repeated_requests_are_throttled(self):
        """The bound on scripting an endpoint that is open by design."""
        self.stub(rows=[])

        with patch.dict(ScopedRateThrottle.THROTTLE_RATES, {"public_export": "2/hour"}):
            first = self.client.get(self.url())
            second = self.client.get(self.url())
            third = self.client.get(self.url())

        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(third.status_code, 429)
